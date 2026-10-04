"""Compare SF2 month-sheet layout across workbook formats.

Usage:
  python tools/sf2-diff/compare.py dump   <file> <sheet> [maxRow]
  python tools/sf2-diff/compare.py diff   <fileA> <sheetA> <fileB> <sheetB> [maxRow]
"""

from __future__ import annotations

import sys
from pathlib import Path


def col_letter(idx: int) -> str:
    """0-based column index -> Excel column letters."""
    s = ""
    idx += 1
    while idx:
        idx, rem = divmod(idx - 1, 26)
        s = chr(65 + rem) + s
    return s


def read_grid(path: Path, sheet: str, max_row: int) -> dict[int, dict[str, str]]:
    """{rowNumber: {columnLetters: value-repr}} of non-empty cells."""
    out: dict[int, dict[str, str]] = {}
    if path.suffix.lower() == ".xls":
        import xlrd

        book = xlrd.open_workbook(str(path), formatting_info=True)
        sh = book.sheet_by_name(sheet)
        for r in range(min(sh.nrows, max_row)):
            row: dict[str, str] = {}
            for c in range(sh.ncols):
                value = sh.cell_value(r, c)
                if value not in ("", None):
                    row[col_letter(c)] = repr(value)
            if row:
                out[r + 1] = row
        return out

    import openpyxl

    book = openpyxl.load_workbook(str(path), data_only=False)
    ws = book[sheet]
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, max_row)):
        cells: dict[str, str] = {}
        for cell in row:
            if cell.value not in ("", None):
                cells[col_letter(cell.column - 1)] = repr(cell.value)
        if cells:
            out[row[0].row] = cells
    return out


def merges_of(path: Path, sheet: str) -> set[str]:
    """Every cell coordinate covered by a merged range."""
    if path.suffix.lower() == ".xls":
        import xlrd

        book = xlrd.open_workbook(str(path), formatting_info=True)
        sh = book.sheet_by_name(sheet)
        return {
            f"{col_letter(c)}{r + 1}"
            for rlo, rhi, clo, chi in sh.merged_cells
            for r in range(rlo, rhi)
            for c in range(clo, chi)
        }

    import openpyxl

    book = openpyxl.load_workbook(str(path))
    ws = book[sheet]
    return {
        f"{col_letter(c - 1)}{r}"
        for rng in ws.merged_cells.ranges
        for r in range(rng.min_row, rng.max_row + 1)
        for c in range(rng.min_col, rng.max_col + 1)
    }


def sorted_cells(cells: set[str], row: int) -> list[str]:
    """The column letters of `cells` that sit on `row`, in column order."""
    prefix = str(row)
    letters = [c[: -len(prefix)] for c in cells if c.endswith(prefix) and c[: -len(prefix)].isalpha()]
    return sorted(letters, key=lambda s: (len(s), s))


def main() -> None:
    mode = sys.argv[1]
    if mode == "dump":
        path, sheet = Path(sys.argv[2]), sys.argv[3]
        max_row = int(sys.argv[4]) if len(sys.argv) > 4 else 40
        for row, cells in read_grid(path, sheet, max_row).items():
            print(f"r{row}: " + " | ".join(f"{c}={v}" for c, v in cells.items()))
        return

    if mode == "diff":
        path_a, sheet_a = Path(sys.argv[2]), sys.argv[3]
        path_b, sheet_b = Path(sys.argv[4]), sys.argv[5]
        max_row = int(sys.argv[6]) if len(sys.argv) > 6 else 40
        grid_a = read_grid(path_a, sheet_a, max_row)
        grid_b = read_grid(path_b, sheet_b, max_row)
        for row in range(1, max_row + 1):
            cells_a = grid_a.get(row, {})
            cells_b = grid_b.get(row, {})
            if cells_a == cells_b:
                continue
            print(f"r{row}:")
            print(f"  A {sheet_a}: " + " | ".join(f"{c}={v}" for c, v in cells_a.items()))
            print(f"  B {sheet_b}: " + " | ".join(f"{c}={v}" for c, v in cells_b.items()))
        return

    if mode == "merges":
        path, sheet = Path(sys.argv[2]), sys.argv[3]
        cells = merges_of(path, sheet)
        for row in (6, 7):
            print(f"row {row} merged cells: {sorted_cells(cells, row)}")
        return

    raise SystemExit(__doc__)


if __name__ == "__main__":
    main()