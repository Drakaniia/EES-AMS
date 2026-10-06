"""Report the ruled lines of an SF2 sheet that do not match the reference form.

The DepEd School Form 2 is a ruled grid, and most of its rules sit on cells that
hold no value at all - the day grid alone borders ~1300 empty cells. A write that
re-merges a range hands every slave cell its master's style, which quietly
replaces the rule drawn *inside* the merge: the sheet keeps the form's text, its
column widths and its horizontal rules, and only the vertical lines come out short
or missing. Nothing a value comparison would notice, so this compares borders.

Usage:
  python tools/sf2-diff/borders.py <file.xlsx> [sheet] --ref <ref.xlsx> [--ref-sheet NAME]
  python tools/sf2-diff/borders.py <file.xlsx> [sheet] --ref <ref.xlsx> --columns=A-AL --rows=1-50

Options:
  --ref <file>      the workbook to compare against
  --ref-sheet NAME  the sheet in that workbook (default: its first sheet)
  --columns=A-AL    restrict the comparison to a column range
  --rows=1-50       restrict the comparison to a row range
  --quiet           print only the summary counts

Every rule is classified as missing (the reference draws it, the sheet does not),
extra (the sheet draws it, the reference does not) or restyled (both draw a rule
there, but not the same one). The exit code is 1 when any rule is missing, so the
script can be used as a check.

Both `--option value` and `--option=value` are accepted.
"""

from __future__ import annotations

import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

# Index 8 is the legacy palette's black - the colour every rule on the form uses.
INDEXED_BLACK = 8

SIDES = ("left", "right", "top", "bottom")


def color_key(color) -> object:
    """A comparable key for a border colour, or None.

    `openpyxl` returns the attribute descriptor rather than a value for whichever
    of `rgb`/`theme`/`indexed` the colour does not use, so the raw `__dict__` is the
    only reliable read.
    """
    if color is None:
        return None
    values = color.__dict__
    kind = values.get("type")
    if kind == "indexed":
        return ("indexed", values.get("indexed"), round(values.get("tint", 0.0), 3))
    if kind == "theme":
        return ("theme", values.get("theme"), round(values.get("tint", 0.0), 3))
    if kind == "rgb":
        return ("rgb", values.get("rgb"), round(values.get("tint", 0.0), 3))
    if kind == "auto":
        return ("auto", None, 0.0)
    return None


def rule(side) -> tuple[str, object] | None:
    """The rule drawn by one side of a cell: `(style, colour)`, or None for no line."""
    if side is None or side.style is None:
        return None
    return (side.style, color_key(side.color))


def describe(value: tuple[str, object] | None) -> str:
    """A readable rendering of one side's rule."""
    if value is None:
        return "none"
    style, color = value
    if color == ("indexed", INDEXED_BLACK, 0.0):
        return style
    return f"{style}/{color}"


def runs(rows: list[int]) -> str:
    """`[8, 9, 10, 12]` -> `8-10, 12`."""
    ordered = sorted(rows)
    spans: list[tuple[int, int]] = []
    start = previous = ordered[0]
    for row in ordered[1:]:
        if row == previous + 1:
            previous = row
            continue
        spans.append((start, previous))
        start = previous = row
    spans.append((start, previous))
    return ", ".join(str(a) if a == b else f"{a}-{b}" for a, b in spans)


def parse_columns(text: str) -> range:
    first, _, last = text.partition("-")
    return range(column_index_from_string(first.upper()), column_index_from_string(last.upper()) + 1)


def parse_rows(text: str) -> range:
    first, _, last = text.partition("-")
    return range(int(first), int(last or first) + 1)


def sheet_of(path: Path, name: str | None):
    workbook = openpyxl.load_workbook(path)
    return workbook[name] if name else workbook.worksheets[0]


def grid(sheet, rows: range, columns: range) -> dict[tuple[int, int], tuple]:
    """`(row, column) -> (left, right, top, bottom)` rules for every cell asked for."""
    out = {}
    for row in rows:
        for column in columns:
            border = sheet.cell(row=row, column=column).border
            out[(row, column)] = (
                rule(border.left),
                rule(border.right),
                rule(border.top),
                rule(border.bottom),
            )
    return out


def report(label: str, found: dict[str, list[tuple[int, int]]], quiet: bool) -> None:
    if quiet or not any(found.values()):
        return
    print(f"\n  -- {label}")
    for side in SIDES:
        cells = found.get(side)
        if not cells:
            continue
        grouped: dict[str, list[int]] = defaultdict(list)
        for row, column in cells:
            grouped[get_column_letter(column)].append(row)
        print(f"     {side}:")
        for column in sorted(grouped, key=column_index_from_string):
            print(f"       {column}: rows {runs(grouped[column])}")


def compare(title: str, target: dict, reference: dict, quiet: bool) -> int:
    missing: dict[str, list[tuple[int, int]]] = defaultdict(list)
    extra: dict[str, list[tuple[int, int]]] = defaultdict(list)
    changed: dict[str, list[tuple[int, int]]] = defaultdict(list)
    examples: dict[str, tuple] = {}

    for key in sorted(target):
        for index, side in enumerate(SIDES):
            want, have = reference[key][index], target[key][index]
            if want == have:
                continue
            if have is None:
                missing[side].append(key)
            elif want is None:
                extra[side].append(key)
            else:
                changed[side].append(key)
            examples.setdefault(side, (key, index, want, have))

    rules = sum(1 for key in target for value in reference[key] if value is not None)
    print(f"\n=== {title}")
    print(f"  {len(target)} cells compared, {rules} rules in the reference")
    print(f"  missing rules : {sum(len(v) for v in missing.values())}")
    print(f"  extra rules   : {sum(len(v) for v in extra.values())}")
    print(f"  restyled rules: {sum(len(v) for v in changed.values())}")

    report("missing rules - the reference draws a line here and the sheet does not", missing, quiet)
    report("extra rules - the sheet draws a line the reference does not", extra, quiet)
    report("restyled rules - a line is drawn, but not the one the reference draws", changed, quiet)

    if not quiet:
        for side, (key, _index, want, have) in examples.items():
            where = f"the {side} edge of {get_column_letter(key[1])}{key[0]}"
            print(f"     e.g. {where}: {describe(have)} (reference {describe(want)})")

    return sum(len(v) for v in missing.values())


def parse_args(argv: list[str]) -> dict:
    """Parse `--option value` and `--option=value` alike."""
    options: dict = {"positional": [], "ref": None, "ref_sheet": None, "quiet": False}
    takes_value = {"--ref": "ref", "--ref-sheet": "ref_sheet", "--rows": "rows", "--columns": "columns"}
    index = 0
    while index < len(argv):
        argument = argv[index]
        name, separator, inline = argument.partition("=")
        if name in takes_value:
            if not separator:
                index += 1
                if index >= len(argv):
                    raise SystemExit(f"{name} needs a value")
                inline = argv[index]
            options[takes_value[name]] = inline
        elif argument == "--quiet":
            options["quiet"] = True
        elif argument in ("--help", "-h"):
            options["help"] = True
        elif argument.startswith("--"):
            raise SystemExit(f"unknown option {argument}\n{__doc__}")
        else:
            options["positional"].append(argument)
        index += 1

    if options.get("rows") is not None:
        options["rows"] = parse_rows(options["rows"])
    if options.get("columns") is not None:
        options["columns"] = parse_columns(options["columns"])
    if options["ref"] is not None:
        options["ref"] = Path(options["ref"])
    return options


def main(argv: list[str]) -> int:
    options = parse_args(argv)
    if options.get("help"):
        print(__doc__)
        return 0

    positional = options["positional"]
    if len(positional) < 1 or options["ref"] is None:
        print(__doc__)
        return 2

    path = Path(positional[0])
    sheet_name = positional[1] if len(positional) > 1 else None
    sheet = sheet_of(path, sheet_name)
    reference = sheet_of(options["ref"], options["ref_sheet"])

    rows = options.get("rows") or range(1, sheet.max_row + 1)
    columns = options.get("columns") or range(1, sheet.max_column + 1)

    title = f"{path.name}!{sheet.title}  vs  {options['ref'].name}!{reference.title}"
    missing = compare(
        title,
        grid(sheet, rows, columns),
        grid(reference, rows, columns),
        options["quiet"],
    )
    return 1 if missing else 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
