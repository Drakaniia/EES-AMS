"""Dump SF2 workbook layout for .xls (xlrd) and .xlsx (openpyxl).

Usage:  python tools/sf2-diff/dump.py <file> [--rows=N] [--sheet=NAME]
Prints: sheet names, merged ranges, and a grid of non-empty cells with coords.
"""

from __future__ import annotations

import sys
from pathlib import Path


def col_letter(idx: int) -> str:
    """0-based column index -> Excel letters."""
    s = ""
    idx += 1
    while idx:
        idx, rem = divmod(idx - 1, 26)
        s = chr(65 + rem) + s
    return s


def dump_xls(path: Path, sheet_name: str | None, max_rows: int) -> None:
    import xlrd

    book = xlrd.open_workbook(str(path), formatting_info=True)
    print(f"FILE {path}  (xls)")
    print("SHEETS:", book.sheet_names())
    for name in book.sheet_names():
        if sheet_name and name != sheet_name:
            continue
        sh = book.sheet_by_name(name)
        print(f"\n=== SHEET {name!r}  dims={sh.nrows}x{sh.ncols} ===")
        print("MERGED:" if False else "MERGED:", [str(r) for r in sh.merged_cells])
        for r in range(min(sh.nrows, max_rows)):
            cells = []
            for c in range(sh.ncols):
                v = sh.cell_value(r, c)
                if v not in ("", None):
                    cells.append(f"{col_letter(c)}{r + 1}={v!r}")
            if cells:
                print(f"r{r + 1}: " + " | ".join(cells))


def dump_xlsx(path: Path, sheet_name: str | None, max_rows: int) -> None:
    import openpyxl

    book = openpyxl.load_workbook(str(path), data_only=False)
    print(f"FILE {path}  (xlsx)")
    print("SHEETS:", book.sheetnames)
    for name in book.sheetnames:
        if sheet_name and name != sheet_name:
            continue
        ws = book[name]
        print(
            f"\n=== SHEET {name!r}  dims={ws.dimensions} "
            f"max_row={ws.max_row} max_col={ws.max_column} ==="
        )
        print("MERGED:", [str(r) for r in ws.merged_cells.ranges])
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, max_rows)):
            cells = []
            for cell in row:
                if cell.value not in ("", None):
                    cells.append(f"{cell.coordinate}={cell.value!r}")
            if cells:
                print(f"r{row[0].row}: " + " | ".join(cells))


def main() -> None:
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    path = Path(args[0])
    sheet_name = None
    max_rows = 60
    for a in sys.argv[1:]:
        if a.startswith("--rows="):
            max_rows = int(a.split("=", 1)[1])
        elif a.startswith("--sheet="):
            sheet_name = a.split("=", 1)[1]
    if path.suffix.lower() == ".xls":
        dump_xls(path, sheet_name, max_rows)
    else:
        dump_xlsx(path, sheet_name, max_rows)


if __name__ == "__main__":
    main()